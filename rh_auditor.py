import os
import re
import unicodedata
from datetime import datetime
import pandas as pd
import numpy as np
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as OpenpyxlImage
import comparator

# ==============================================================================
# CORES E ESTILOS PADRONIZADOS DA IDENTIDADE VISUAL VALIDAÍ
# ==============================================================================
COLOR_GREEN = "4FD580"
COLOR_DARK = "0F172A"
COLOR_GRAY_BG = "F8FAFC"
COLOR_SUCCESS_BG = "DCFCE7"
COLOR_SUCCESS_TEXT = "15803D"
COLOR_ERROR_BG = "FEE2E2"
COLOR_ERROR_TEXT = "B91C1C"
COLOR_WARN_BG = "FEF3C7"
COLOR_WARN_TEXT = "B45309"
COLOR_ORANGE_BG = "FFEDD5"
COLOR_ORANGE_TEXT = "C2410C"
COLOR_BLUE_BG = "E0F2FE"
COLOR_BLUE_TEXT = "0369A1"


# ==============================================================================
# 1. UTILITÁRIOS DE NORMALIZAÇÃO E LIMPEZA
# ==============================================================================

def clean_digits(val):
    """Extrai apenas dígitos numéricos de uma string ou número."""
    if val is None or pd.isna(val):
        return ""
    s = str(val).strip()
    if s.endswith('.0'):
        s = s[:-2]
    return re.sub(r'\D', '', s)


def normalize_cpf_or_key(val, is_cpf=True):
    """
    Normaliza a chave do colaborador (CPF, Matrícula, ID).
    Se for CPF (ou tiver de 8 a 11 dígitos), remove pontuação e aplica .zfill(11)
    para recuperar zeros à esquerda engolidos pelo Excel e evitar falsos órfãos.
    """
    if val is None or pd.isna(val):
        return ""
    s = str(val).strip()
    if s.endswith('.0'):
        s = s[:-2]

    digits = re.sub(r'\D', '', s)
    if is_cpf or (8 <= len(digits) <= 11 and digits.isdigit()):
        return digits.zfill(11)
    elif digits:
        return digits
    return s.strip().upper()


def normalize_name_string(val):
    """Normaliza texto/nome: sem acentos, sem espaços extras e em caixa alta."""
    if val is None or pd.isna(val):
        return ""
    s = str(val).strip()
    s = re.sub(r'\s+', ' ', s)
    s = unicodedata.normalize('NFKD', s).encode('ASCII', 'ignore').decode('utf-8')
    return s.upper()


def parse_currency(val):
    """Converte qualquer formato monetário em float."""
    return comparator.parse_financial_value(val)


def format_brl(val):
    """Formata número para moeda brasileira R$ #.##0,00."""
    if val is None or pd.isna(val):
        return "R$ 0,00"
    try:
        v = float(val)
        neg = "-" if v < 0 else ""
        v = abs(v)
        s = f"{v:,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.')
        return f"{neg}R$ {s}"
    except Exception:
        return "R$ 0,00"


def format_pct(val):
    """Formata número para porcentagem com 2 casas decimais."""
    if val is None or pd.isna(val):
        return "0,00%"
    try:
        v = float(val)
        sign = "+" if v > 0 else ""
        return f"{sign}{v:.2f}%".replace('.', ',')
    except Exception:
        return "0,00%"


def is_likely_cpf_column(col_name):
    """Verifica se o nome da coluna remete a CPF."""
    c = comparator.normalize_column_name(col_name)
    return any(k in c for k in ['CPF', 'DOC', 'DOCUMENTO', 'CIC'])


def is_likely_employee_key(col_name):
    """Verifica se o nome da coluna remete a identificador de colaborador."""
    c = comparator.normalize_column_name(col_name)
    return any(k in c for k in [
        'CPF', 'MATRICULA', 'ID_FUNCIONARIO', 'ID_COLABORADOR',
        'PIS', 'REGISTRO', 'CODIGO_FUNC', 'CHAVE'
    ])


def suggest_employee_keys(cols_ref, cols_audit):
    """Sugere a melhor chave primária e complementar para cruzamento de RH."""
    # 1. Tenta CPF primeiro
    cpf_ref = next((c for c in cols_ref if is_likely_cpf_column(c)), None)
    cpf_audit = next((c for c in cols_audit if is_likely_cpf_column(c)), None)

    if cpf_ref and cpf_audit:
        return {
            'primary_ref': cpf_ref,
            'primary_audit': cpf_audit,
            'is_cpf': True
        }

    # 2. Tenta Matrícula / ID
    key_ref = next((c for c in cols_ref if is_likely_employee_key(c)), None)
    key_audit = next((c for c in cols_audit if is_likely_employee_key(c)), None)

    if key_ref and key_audit:
        return {
            'primary_ref': key_ref,
            'primary_audit': key_audit,
            'is_cpf': is_likely_cpf_column(key_ref) or is_likely_cpf_column(key_audit)
        }

    # 3. Fallback: primeira coluna comum
    common = [c for c in cols_ref if c in cols_audit]
    if common:
        return {
            'primary_ref': common[0],
            'primary_audit': common[0],
            'is_cpf': is_likely_cpf_column(common[0])
        }

    return {
        'primary_ref': cols_ref[0] if cols_ref else '',
        'primary_audit': cols_audit[0] if cols_audit else '',
        'is_cpf': False
    }


def find_name_column(columns):
    """Identifica coluna com o nome do colaborador."""
    name_keywords = ['NOME', 'COLABORADOR', 'FUNCIONARIO', 'TITULAR', 'EMPREGADO', 'NM_']
    for col in columns:
        col_norm = comparator.normalize_column_name(col)
        if any(k in col_norm for k in name_keywords):
            return col
    return None


def is_numeric_series(series, sample_size=50):
    """Detecta se uma série de dados é composta por valores monetários ou numéricos."""
    valid_count = 0
    numeric_count = 0
    sample = series.dropna().head(sample_size)
    for v in sample:
        s = str(v).strip()
        if not s:
            continue
        valid_count += 1
        clean_s = s.replace('R$', '').replace('$', '').replace(' ', '').replace('.', '').replace(',', '.')
        try:
            float(clean_s)
            numeric_count += 1
        except ValueError:
            pass
    if valid_count == 0:
        return False
    return (numeric_count / valid_count) >= 0.7


# ==============================================================================
# 2. MOTOR PRINCIPAL DE AUDITORIA DE FOLHA E RH
# ==============================================================================

def audit_rh_bases(df_ref: pd.DataFrame, df_audit: pd.DataFrame, options: dict) -> dict:
    """
    Executa a auditoria completa de duas bases de RH:
    - Modo Folha: Mês Anterior vs. Mês Atual
    - Modo Benefícios: Ativos RH vs. Fatura da Operadora
    """
    mode = options.get('mode', 'folha')
    key_ref_col = options.get('key_ref')
    key_audit_col = options.get('key_audit')
    sec_key_ref = options.get('sec_key_ref')
    sec_key_audit = options.get('sec_key_audit')
    normalize_cpf = bool(options.get('normalize_cpf', True))
    ignore_accents = bool(options.get('ignore_accents', True))
    variation_threshold = float(options.get('variation_threshold', 10.0))
    column_mapping = options.get('column_mapping', {})

    if not key_ref_col or key_ref_col not in df_ref.columns:
        raise ValueError(f"Coluna chave da Base de Referência '{key_ref_col}' não encontrada.")
    if not key_audit_col or key_audit_col not in df_audit.columns:
        raise ValueError(f"Coluna chave da Base Auditada '{key_audit_col}' não encontrada.")

    # Detecção se a chave primária é CPF
    is_cpf_ref = is_likely_cpf_column(key_ref_col)
    is_cpf_audit = is_likely_cpf_column(key_audit_col)
    is_cpf = normalize_cpf and (is_cpf_ref or is_cpf_audit)

    # Identificação de colunas de Nome para exibição amigável
    name_col_ref = options.get('name_col_ref') or find_name_column(df_ref.columns)
    name_col_audit = options.get('name_col_audit') or find_name_column(df_audit.columns)

    # Preparar cópias para normalização
    df_ref_clean = df_ref.copy()
    df_audit_clean = df_audit.copy()

    # Criar chaves normalizadas
    def build_norm_key(row, key_col, sec_col):
        val = row.get(key_col, '')
        k = normalize_cpf_or_key(val, is_cpf=is_cpf)
        if sec_col and sec_col in row and pd.notna(row[sec_col]):
            sec_val = str(row[sec_col]).strip()
            if ignore_accents:
                sec_val = normalize_name_string(sec_val)
            k = f"{k}__{sec_val}"
        return k

    df_ref_clean['_AUDIT_KEY_'] = df_ref_clean.apply(
        lambda r: build_norm_key(r, key_ref_col, sec_key_ref), axis=1
    )
    df_audit_clean['_AUDIT_KEY_'] = df_audit_clean.apply(
        lambda r: build_norm_key(r, key_audit_col, sec_key_audit), axis=1
    )

    # Remover linhas com chave vazia
    df_ref_clean = df_ref_clean[df_ref_clean['_AUDIT_KEY_'] != ''].drop_duplicates(subset=['_AUDIT_KEY_'], keep='first')
    df_audit_clean = df_audit_clean[df_audit_clean['_AUDIT_KEY_'] != ''].drop_duplicates(subset=['_AUDIT_KEY_'], keep='first')

    keys_ref_set = set(df_ref_clean['_AUDIT_KEY_'])
    keys_audit_set = set(df_audit_clean['_AUDIT_KEY_'])

    common_keys = sorted(list(keys_ref_set.intersection(keys_audit_set)))
    exclusive_ref_keys = sorted(list(keys_ref_set - keys_audit_set))
    exclusive_audit_keys = sorted(list(keys_audit_set - keys_ref_set))

    # Mapeamento de colunas numéricas
    numeric_cols_map = {}
    for col_r, col_a in column_mapping.items():
        if col_r in df_ref.columns and col_a in df_audit.columns:
            if is_numeric_series(df_ref[col_r]) or is_numeric_series(df_audit[col_a]):
                numeric_cols_map[col_r] = col_a

    # Dicionários de acesso rápido indexados por _AUDIT_KEY_
    ref_dict = df_ref_clean.set_index('_AUDIT_KEY_').to_dict(orient='index')
    audit_dict = df_audit_clean.set_index('_AUDIT_KEY_').to_dict(orient='index')

    # Estruturas de resultado
    divergences_list = []
    critical_variations_list = []
    employees_with_divergence = set()

    # 1. Análise detalhada dos colaboradores comuns (cruzamento)
    for k in common_keys:
        row_ref = ref_dict[k]
        row_audit = audit_dict[k]

        emp_name = ""
        if name_col_ref and name_col_ref in row_ref and pd.notna(row_ref[name_col_ref]):
            emp_name = str(row_ref[name_col_ref]).strip()
        elif name_col_audit and name_col_audit in row_audit and pd.notna(row_audit[name_col_audit]):
            emp_name = str(row_audit[name_col_audit]).strip()

        emp_code = str(row_ref.get(key_ref_col, k)).strip()
        if is_cpf:
            emp_code = normalize_cpf_or_key(emp_code, is_cpf=True)

        for col_r, col_a in column_mapping.items():
            if col_r not in row_ref or col_a not in row_audit:
                continue

            val_r = row_ref[col_r]
            val_a = row_audit[col_a]

            is_numeric = (col_r in numeric_cols_map)

            if is_numeric:
                num_r = parse_currency(val_r)
                num_a = parse_currency(val_a)
                diff = round(num_a - num_r, 2)

                if abs(diff) >= 0.01:
                    employees_with_divergence.add(k)

                    # Cálculo percentual
                    if num_r != 0:
                        var_pct = round(((num_a - num_r) / abs(num_r)) * 100.0, 2)
                    elif num_a != 0:
                        var_pct = 100.0
                    else:
                        var_pct = 0.0

                    is_critical = abs(var_pct) >= variation_threshold

                    status = "Aumento Atípico" if (var_pct > 0 and is_critical) else (
                        "Queda Atípica" if (var_pct < 0 and is_critical) else (
                            "Aumento" if var_pct > 0 else "Redução"
                        )
                    )

                    item = {
                        'employee_key': emp_code,
                        'employee_name': emp_name,
                        'col_ref': col_r,
                        'col_audit': col_a,
                        'col_display': f"{col_r} → {col_a}" if col_r != col_a else col_r,
                        'val_ref': num_r,
                        'val_ref_fmt': format_brl(num_r),
                        'val_audit': num_a,
                        'val_audit_fmt': format_brl(num_a),
                        'diff': diff,
                        'diff_fmt': format_brl(diff),
                        'var_pct': var_pct,
                        'var_pct_fmt': format_pct(var_pct),
                        'is_numeric': True,
                        'is_critical': is_critical,
                        'status': status
                    }
                    divergences_list.append(item)
                    if is_critical:
                        critical_variations_list.append(item)

            else:
                # Campo textual ou cadastral (Cargo, Departamento, etc.)
                str_r = str(val_r).strip() if pd.notna(val_r) else ""
                str_a = str(val_a).strip() if pd.notna(val_a) else ""

                cmp_r = normalize_name_string(str_r) if ignore_accents else str_r.upper()
                cmp_a = normalize_name_string(str_a) if ignore_accents else str_a.upper()

                if cmp_r != cmp_a:
                    employees_with_divergence.add(k)
                    item = {
                        'employee_key': emp_code,
                        'employee_name': emp_name,
                        'col_ref': col_r,
                        'col_audit': col_a,
                        'col_display': f"{col_r} → {col_a}" if col_r != col_a else col_r,
                        'val_ref': str_r,
                        'val_ref_fmt': str_r or "(Vazio)",
                        'val_audit': str_a,
                        'val_audit_fmt': str_a or "(Vazio)",
                        'diff': 0.0,
                        'diff_fmt': "-",
                        'var_pct': 0.0,
                        'var_pct_fmt': "Alteração Cadastral",
                        'is_numeric': False,
                        'is_critical': False,
                        'status': "Alteração Cadastral"
                    }
                    divergences_list.append(item)

    # 2. Lista de Exclusivos Base 1 (Possíveis Desligados / Não Faturados)
    exclusive_ref_list = []
    for k in exclusive_ref_keys:
        row = ref_dict[k]
        emp_name = str(row.get(name_col_ref, '')).strip() if name_col_ref else ""
        emp_code = str(row.get(key_ref_col, k)).strip()
        if is_cpf:
            emp_code = normalize_cpf_or_key(emp_code, is_cpf=True)

        # Busca valores monetários relevantes (ex: Salário Base ou Líquido)
        sal_val = 0.0
        cargo = ""
        depto = ""
        for c in df_ref.columns:
            cnorm = comparator.normalize_column_name(c)
            if any(w in cnorm for w in ['SALARIO', 'LIQUIDO', 'BRUTO', 'VALOR']):
                if sal_val == 0.0:
                    sal_val = parse_currency(row.get(c, 0.0))
            if 'CARGO' in cnorm or 'FUNCAO' in cnorm:
                cargo = str(row.get(c, '')).strip()
            if 'DEP' in cnorm or 'SETOR' in cnorm:
                depto = str(row.get(c, '')).strip()

        status_label = "Desligado / Saída da Folha" if mode == 'folha' else "Ativo sem Fatura no Benefício"

        exclusive_ref_list.append({
            'employee_key': emp_code,
            'employee_name': emp_name,
            'cargo': cargo or "-",
            'departamento': depto or "-",
            'valor_referencia': sal_val,
            'valor_referencia_fmt': format_brl(sal_val),
            'status': status_label
        })

    # 3. Lista de Exclusivos Base 2 (Admitidos / Cobrança Indevida / Funcionários Fantasmas)
    exclusive_audit_list = []
    for k in exclusive_audit_keys:
        row = audit_dict[k]
        emp_name = str(row.get(name_col_audit, '')).strip() if name_col_audit else ""
        emp_code = str(row.get(key_audit_col, k)).strip()
        if is_cpf:
            emp_code = normalize_cpf_or_key(emp_code, is_cpf=True)

        sal_val = 0.0
        cargo = ""
        depto = ""
        for c in df_audit.columns:
            cnorm = comparator.normalize_column_name(c)
            if any(w in cnorm for w in ['SALARIO', 'LIQUIDO', 'BRUTO', 'VALOR', 'MENSALIDADE']):
                if sal_val == 0.0:
                    sal_val = parse_currency(row.get(c, 0.0))
            if 'CARGO' in cnorm or 'PLANO' in cnorm:
                cargo = str(row.get(c, '')).strip()
            if 'DEP' in cnorm or 'SETOR' in cnorm:
                depto = str(row.get(c, '')).strip()

        status_label = "Novo Admitido na Folha" if mode == 'folha' else "Cobrança Indevida / Funcionário Fantasma"

        exclusive_audit_list.append({
            'employee_key': emp_code,
            'employee_name': emp_name,
            'cargo': cargo or "-",
            'departamento': depto or "-",
            'valor_auditado': sal_val,
            'valor_auditado_fmt': format_brl(sal_val),
            'status': status_label
        })

    # 4. Sumarização Financeira por Coluna Mapeada
    financial_summary = []
    total_net_ref_sum = 0.0
    total_net_audit_sum = 0.0

    for col_r, col_a in numeric_cols_map.items():
        sum_r = float(df_ref[col_r].apply(parse_currency).sum())
        sum_a = float(df_audit[col_a].apply(parse_currency).sum())
        diff = round(float(sum_a - sum_r), 2)
        pct = round(float(((sum_a - sum_r) / sum_r) * 100.0), 2) if sum_r != 0 else 0.0

        if pct == 0.0 and abs(diff) < 0.01:
            situacao = "Sem Variação"
        elif diff > 0:
            situacao = "Aumento Financeiro"
        else:
            situacao = "Redução Financeira"

        financial_summary.append({
            'col_ref': col_r,
            'col_audit': col_a,
            'col_display': f"{col_r} vs {col_a}",
            'total_ref': float(sum_r),
            'total_ref_fmt': format_brl(sum_r),
            'total_audit': float(sum_a),
            'total_audit_fmt': format_brl(sum_a),
            'diff': float(diff),
            'diff_fmt': format_brl(diff),
            'var_pct': float(pct),
            'var_pct_fmt': format_pct(pct),
            'situacao': situacao
        })

        col_norm = comparator.normalize_column_name(col_r)
        if any(w in col_norm for w in ['LIQUIDO', 'BRUTO', 'VALOR_TOTAL', 'TOTAL', 'MENSALIDADE']):
            if total_net_ref_sum == 0.0:
                total_net_ref_sum = float(sum_r)
                total_net_audit_sum = float(sum_a)

    # Fallback para impacto financeiro total se nenhuma coluna tiver nome explícito
    if total_net_ref_sum == 0.0 and financial_summary:
        total_net_ref_sum = float(financial_summary[0]['total_ref'])
        total_net_audit_sum = float(financial_summary[0]['total_audit'])

    total_financial_impact = round(float(total_net_audit_sum - total_net_ref_sum), 2)

    # 5. Cálculo dos 6 KPIs principais
    total_audited_count = int(len(common_keys))
    divergence_count = int(len(employees_with_divergence))
    divergence_pct = round(float((divergence_count / total_audited_count) * 100.0), 1) if total_audited_count > 0 else 0.0

    kpis = {
        'total_audited': total_audited_count,
        'divergence_count': divergence_count,
        'divergence_pct': float(divergence_pct),
        'critical_variations_count': int(len(critical_variations_list)),
        'variation_threshold_pct': float(variation_threshold),
        'exclusive_ref_count': int(len(exclusive_ref_list)),
        'exclusive_audit_count': int(len(exclusive_audit_list)),
        'financial_impact': float(total_financial_impact),
        'financial_impact_fmt': format_brl(total_financial_impact),
        'total_ref_sum': float(total_net_ref_sum),
        'total_ref_sum_fmt': format_brl(total_net_ref_sum),
        'total_audit_sum': float(total_net_audit_sum),
        'total_audit_sum_fmt': format_brl(total_net_audit_sum),
        'mode': mode
    }

    # Ordenações
    critical_variations_list.sort(key=lambda x: abs(x['diff']), reverse=True)
    divergences_list.sort(key=lambda x: abs(x['diff']), reverse=True)
    exclusive_ref_list.sort(key=lambda x: x['valor_referencia'], reverse=True)
    exclusive_audit_list.sort(key=lambda x: x['valor_auditado'], reverse=True)

    return {
        'kpis': kpis,
        'financial_summary': financial_summary,
        'divergences': divergences_list,
        'critical_variations': critical_variations_list,
        'exclusive_ref': exclusive_ref_list,
        'exclusive_audit': exclusive_audit_list,
        'mode': mode
    }


# ==============================================================================
# 3. GERADOR DE RELATÓRIO EXCEL COMPLETO DE AUDITORIA DE RH (5 ABAS)
# ==============================================================================

def generate_rh_audit_excel(audit_data: dict, output_path: str, logo_path: str = None) -> str:
    """
    Gera arquivo Excel (.xlsx) profissional no padrão visual ValidAí com 5 abas:
    - Aba 1: Resumo_RH (KPIs executivos, Sumarização Financeira e Logo)
    - Aba 2: Divergencias_Detalhadas (Todas as diferenças cadastrais e monetárias)
    - Aba 3: Variacoes_Acima_Limite (Apenas variações atípicas >= X%)
    - Aba 4: Desligados_Apenas_Base1 (Colaboradores presentes só na Base 1)
    - Aba 5: Admitidos_Ou_Indevidos_Base2 (Colaboradores presentes só na Base 2)
    """
    kpis = audit_data['kpis']
    fin_summary = audit_data['financial_summary']
    divergences = audit_data['divergences']
    critical_vars = audit_data['critical_variations']
    exclusive_ref = audit_data['exclusive_ref']
    exclusive_audit = audit_data['exclusive_audit']
    mode = audit_data.get('mode', 'folha')

    wb = openpyxl.Workbook()
    # Remove aba padrão inicial
    default_sheet = wb.active
    wb.remove(default_sheet)

    font_family = "Calibri"
    font_title = Font(name=font_family, size=15, bold=True, color=COLOR_DARK)
    font_sub = Font(name=font_family, size=10, italic=True, color="64748B")
    font_sec = Font(name=font_family, size=11, bold=True, color="0F172A")
    font_bold = Font(name=font_family, size=11, bold=True, color=COLOR_DARK)
    font_normal = Font(name=font_family, size=11, color=COLOR_DARK)
    font_header = Font(name=font_family, size=11, bold=True, color="FFFFFF")

    fill_header = PatternFill(start_color=COLOR_GREEN, end_color=COLOR_GREEN, fill_type="solid")
    fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    fill_ok = PatternFill(start_color=COLOR_SUCCESS_BG, end_color=COLOR_SUCCESS_BG, fill_type="solid")
    font_ok = Font(name=font_family, size=10, bold=True, color=COLOR_SUCCESS_TEXT)
    fill_error = PatternFill(start_color=COLOR_ERROR_BG, end_color=COLOR_ERROR_BG, fill_type="solid")
    font_error = Font(name=font_family, size=10, bold=True, color=COLOR_ERROR_TEXT)
    fill_warn = PatternFill(start_color=COLOR_WARN_BG, end_color=COLOR_WARN_BG, fill_type="solid")
    font_warn = Font(name=font_family, size=10, bold=True, color=COLOR_WARN_TEXT)
    fill_blue = PatternFill(start_color=COLOR_BLUE_BG, end_color=COLOR_BLUE_BG, fill_type="solid")
    font_blue = Font(name=font_family, size=10, bold=True, color=COLOR_BLUE_TEXT)

    thin_border = Border(
        left=Side(style='thin', color="CBD5E1"),
        right=Side(style='thin', color="CBD5E1"),
        top=Side(style='thin', color="CBD5E1"),
        bottom=Side(style='thin', color="CBD5E1")
    )

    currency_format = 'R$ #,##0.00'
    pct_format = '0.00%'

    # ---------------------------------------------------------
    # ABA 1: Resumo_RH
    # ---------------------------------------------------------
    ws_resumo = wb.create_sheet(title="Resumo_RH")
    ws_resumo.views.sheetView[0].showGridLines = True

    # Inserção do Logo ValidAí
    if logo_path and os.path.exists(logo_path):
        try:
            img = OpenpyxlImage(logo_path)
            img.width = 135
            img.height = 38
            ws_resumo.add_image(img, "B2")
        except Exception:
            pass

    ws_resumo["D2"] = "RELATÓRIO DE AUDITORIA DE FOLHA E RH"
    ws_resumo["D2"].font = font_title
    mode_label = "Comparativo Mensal de Folha de Pagamento" if mode == 'folha' else "Ativos no RH vs. Fatura de Benefícios"
    ws_resumo["D3"] = f"Módulo Especializado ValidAí | {mode_label} | Gerado em {datetime.now().strftime('%d/%m/%Y às %H:%M')}"
    ws_resumo["D3"].font = font_sub

    # Bloco dos 6 KPIs
    ws_resumo["B6"] = "INDICADORES EXECUTIVOS DA AUDITORIA"
    ws_resumo["B6"].font = font_sec

    kpi_rows = [
        ("Colaboradores Auditados (Cruzados)", kpis['total_audited'], "Total de colaboradores presentes em ambas as bases"),
        ("Colaboradores com Divergência", f"{kpis['divergence_count']} ({kpis['divergence_pct']}%)", "Cadastros com alguma alteração cadastral ou de valor"),
        (f"Alertas de Variação Atípica (> {kpis['variation_threshold_pct']}%)", kpis['critical_variations_count'], "Salários ou valores com oscilação acima da tolerância"),
        ("Exclusivos Base 1 (Possíveis Desligados)", kpis['exclusive_ref_count'], "Colaboradores presentes apenas na base de referência"),
        ("Exclusivos Base 2 (Admitidos / Cobrança Indevida)", kpis['exclusive_audit_count'], "Colaboradores novos ou cobranças sem cadastro ativo"),
        ("Impacto Financeiro Líquido", kpis['financial_impact_fmt'], "Diferença líquida total (Base Auditada - Base Referência)")
    ]

    r_idx = 7
    ws_resumo.cell(row=r_idx, column=2, value="Indicador").font = font_header
    ws_resumo.cell(row=r_idx, column=2).fill = fill_header
    ws_resumo.cell(row=r_idx, column=2).border = thin_border

    ws_resumo.cell(row=r_idx, column=3, value="Valor").font = font_header
    ws_resumo.cell(row=r_idx, column=3).fill = fill_header
    ws_resumo.cell(row=r_idx, column=3).alignment = Alignment(horizontal="center")
    ws_resumo.cell(row=r_idx, column=3).border = thin_border

    ws_resumo.cell(row=r_idx, column=4, value="Descrição / Interpretação").font = font_header
    ws_resumo.cell(row=r_idx, column=4).fill = fill_header
    ws_resumo.cell(row=r_idx, column=4).border = thin_border

    for item in kpi_rows:
        r_idx += 1
        c2 = ws_resumo.cell(row=r_idx, column=2, value=item[0])
        c2.font = font_bold
        c2.border = thin_border
        c2.fill = fill_zebra

        c3 = ws_resumo.cell(row=r_idx, column=3, value=item[1])
        c3.font = font_bold
        c3.alignment = Alignment(horizontal="center")
        c3.border = thin_border

        c4 = ws_resumo.cell(row=r_idx, column=4, value=item[2])
        c4.font = font_normal
        c4.border = thin_border

    # Tabela de Sumarização Financeira
    r_idx += 3
    ws_resumo.cell(row=r_idx, column=2, value="SUMARIZAÇÃO FINANCEIRA DA FOLHA / BENEFÍCIOS").font = font_sec

    r_idx += 1
    fin_headers = [
        "Coluna Base 1", "Coluna Base 2", "Total Anterior (R$)",
        "Total Atual (R$)", "Diferença Líquida (R$)", "% Variação", "Situação"
    ]
    for c_idx, h_text in enumerate(fin_headers, start=2):
        cell = ws_resumo.cell(row=r_idx, column=c_idx, value=h_text)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center" if c_idx >= 4 else "left")
        cell.border = thin_border

    for f_item in fin_summary:
        r_idx += 1
        ws_resumo.cell(row=r_idx, column=2, value=f_item['col_ref']).font = font_normal
        ws_resumo.cell(row=r_idx, column=3, value=f_item['col_audit']).font = font_normal

        c_t1 = ws_resumo.cell(row=r_idx, column=4, value=f_item['total_ref'])
        c_t1.number_format = currency_format
        c_t1.alignment = Alignment(horizontal="right")

        c_t2 = ws_resumo.cell(row=r_idx, column=5, value=f_item['total_audit'])
        c_t2.number_format = currency_format
        c_t2.alignment = Alignment(horizontal="right")

        c_diff = ws_resumo.cell(row=r_idx, column=6, value=f_item['diff'])
        c_diff.number_format = currency_format
        c_diff.alignment = Alignment(horizontal="right")

        c_pct = ws_resumo.cell(row=r_idx, column=7, value=f_item['var_pct'] / 100.0)
        c_pct.number_format = pct_format
        c_pct.alignment = Alignment(horizontal="center")

        c_sit = ws_resumo.cell(row=r_idx, column=8, value=f_item['situacao'])
        c_sit.alignment = Alignment(horizontal="center")
        if f_item['diff'] > 0:
            c_sit.fill = fill_warn
            c_sit.font = font_warn
        elif f_item['diff'] < 0:
            c_sit.fill = fill_blue
            c_sit.font = font_blue
        else:
            c_sit.fill = fill_ok
            c_sit.font = font_ok

        for col_b in range(2, 9):
            ws_resumo.cell(row=r_idx, column=col_b).border = thin_border

    # ---------------------------------------------------------
    # ABA 2: Divergencias_Detalhadas
    # ---------------------------------------------------------
    ws_div = wb.create_sheet(title="Divergencias_Detalhadas")
    ws_div.views.sheetView[0].showGridLines = True

    div_headers = [
        "CPF / Matrícula", "Nome do Colaborador", "Coluna", "Valor Anterior (Base 1)",
        "Valor Atual (Base 2)", "Diferença R$", "% Variação", "Status / Gravidade"
    ]
    for c_idx, h_text in enumerate(div_headers, start=1):
        cell = ws_div.cell(row=1, column=c_idx, value=h_text)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center" if c_idx >= 4 else "left")
        cell.border = thin_border

    for r_i, d in enumerate(divergences, start=2):
        ws_div.cell(row=r_i, column=1, value=d['employee_key']).font = font_bold
        ws_div.cell(row=r_i, column=2, value=d['employee_name']).font = font_normal
        ws_div.cell(row=r_i, column=3, value=d['col_display']).font = font_normal

        c_val1 = ws_div.cell(row=r_i, column=4, value=d['val_ref'] if d['is_numeric'] else str(d['val_ref']))
        c_val2 = ws_div.cell(row=r_i, column=5, value=d['val_audit'] if d['is_numeric'] else str(d['val_audit']))

        if d['is_numeric']:
            c_val1.number_format = currency_format
            c_val1.alignment = Alignment(horizontal="right")
            c_val2.number_format = currency_format
            c_val2.alignment = Alignment(horizontal="right")

            c_diff = ws_div.cell(row=r_i, column=6, value=d['diff'])
            c_diff.number_format = currency_format
            c_diff.alignment = Alignment(horizontal="right")

            c_pct = ws_div.cell(row=r_i, column=7, value=d['var_pct'] / 100.0)
            c_pct.number_format = pct_format
            c_pct.alignment = Alignment(horizontal="center")
        else:
            ws_div.cell(row=r_i, column=6, value="-").alignment = Alignment(horizontal="center")
            ws_div.cell(row=r_i, column=7, value="Cadastral").alignment = Alignment(horizontal="center")

        c_st = ws_div.cell(row=r_i, column=8, value=d['status'])
        c_st.alignment = Alignment(horizontal="center")
        if d.get('is_critical'):
            c_st.fill = fill_error
            c_st.font = font_error
        elif d['is_numeric']:
            c_st.fill = fill_warn
            c_st.font = font_warn
        else:
            c_st.fill = fill_blue
            c_st.font = font_blue

        for col_b in range(1, 9):
            ws_div.cell(row=r_i, column=col_b).border = thin_border
            if r_i % 2 == 1 and not d.get('is_critical'):
                ws_div.cell(row=r_i, column=col_b).fill = fill_zebra

    # ---------------------------------------------------------
    # ABA 3: Variacoes_Acima_Limite
    # ---------------------------------------------------------
    ws_crit = wb.create_sheet(title="Variacoes_Acima_Limite")
    ws_crit.views.sheetView[0].showGridLines = True

    crit_headers = [
        "CPF / Matrícula", "Nome do Colaborador", "Coluna Auditada", "Valor Anterior",
        "Valor Atual", "Diferença Líquida (R$)", "% Variação", "Alerta Crítico"
    ]
    for c_idx, h_text in enumerate(crit_headers, start=1):
        cell = ws_crit.cell(row=1, column=c_idx, value=h_text)
        cell.font = font_header
        cell.fill = PatternFill(start_color="DC2626", end_color="DC2626", fill_type="solid")
        cell.alignment = Alignment(horizontal="center" if c_idx >= 4 else "left")
        cell.border = thin_border

    for r_i, c in enumerate(critical_vars, start=2):
        ws_crit.cell(row=r_i, column=1, value=c['employee_key']).font = font_bold
        ws_crit.cell(row=r_i, column=2, value=c['employee_name']).font = font_normal
        ws_crit.cell(row=r_i, column=3, value=c['col_display']).font = font_normal

        c1 = ws_crit.cell(row=r_i, column=4, value=c['val_ref'])
        c1.number_format = currency_format
        c1.alignment = Alignment(horizontal="right")

        c2 = ws_crit.cell(row=r_i, column=5, value=c['val_audit'])
        c2.number_format = currency_format
        c2.alignment = Alignment(horizontal="right")

        c_df = ws_crit.cell(row=r_i, column=6, value=c['diff'])
        c_df.number_format = currency_format
        c_df.alignment = Alignment(horizontal="right")

        c_pt = ws_crit.cell(row=r_i, column=7, value=c['var_pct'] / 100.0)
        c_pt.number_format = pct_format
        c_pt.alignment = Alignment(horizontal="center")

        st_c = ws_crit.cell(row=r_i, column=8, value=f"Alerta: {c['status']} ({c['var_pct_fmt']})")
        st_c.fill = fill_error
        st_c.font = font_error
        st_c.alignment = Alignment(horizontal="center")

        for col_b in range(1, 9):
            ws_crit.cell(row=r_i, column=col_b).border = thin_border

    # ---------------------------------------------------------
    # ABA 4: Desligados_Apenas_Base1
    # ---------------------------------------------------------
    ws_ex1 = wb.create_sheet(title="Desligados_Apenas_Base1")
    ws_ex1.views.sheetView[0].showGridLines = True

    ex1_headers = [
        "CPF / Matrícula", "Nome do Colaborador", "Cargo / Função",
        "Departamento", "Último Valor Registrado (R$)", "Situação Identificada"
    ]
    for c_idx, h_text in enumerate(ex1_headers, start=1):
        cell = ws_ex1.cell(row=1, column=c_idx, value=h_text)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center" if c_idx >= 5 else "left")
        cell.border = thin_border

    for r_i, e in enumerate(exclusive_ref, start=2):
        ws_ex1.cell(row=r_i, column=1, value=e['employee_key']).font = font_bold
        ws_ex1.cell(row=r_i, column=2, value=e['employee_name']).font = font_normal
        ws_ex1.cell(row=r_i, column=3, value=e['cargo']).font = font_normal
        ws_ex1.cell(row=r_i, column=4, value=e['departamento']).font = font_normal

        c_val = ws_ex1.cell(row=r_i, column=5, value=e['valor_referencia'])
        c_val.number_format = currency_format
        c_val.alignment = Alignment(horizontal="right")

        c_st = ws_ex1.cell(row=r_i, column=6, value=e['status'])
        c_st.fill = fill_warn
        c_st.font = font_warn
        c_st.alignment = Alignment(horizontal="center")

        for col_b in range(1, 7):
            ws_ex1.cell(row=r_i, column=col_b).border = thin_border
            if r_i % 2 == 1:
                ws_ex1.cell(row=r_i, column=col_b).fill = fill_zebra

    # ---------------------------------------------------------
    # ABA 5: Admitidos_Ou_Indevidos_Base2
    # ---------------------------------------------------------
    ws_ex2 = wb.create_sheet(title="Admitidos_Ou_Indevidos_Base2")
    ws_ex2.views.sheetView[0].showGridLines = True

    ex2_headers = [
        "CPF / Matrícula", "Nome do Colaborador", "Cargo / Plano",
        "Departamento", "Valor Faturado / Salário (R$)", "Situação / Alerta de Cobrança"
    ]
    for c_idx, h_text in enumerate(ex2_headers, start=1):
        cell = ws_ex2.cell(row=1, column=c_idx, value=h_text)
        cell.font = font_header
        cell.fill = PatternFill(start_color="F59E0B", end_color="F59E0B", fill_type="solid")
        cell.alignment = Alignment(horizontal="center" if c_idx >= 5 else "left")
        cell.border = thin_border

    for r_i, e in enumerate(exclusive_audit, start=2):
        ws_ex2.cell(row=r_i, column=1, value=e['employee_key']).font = font_bold
        ws_ex2.cell(row=r_i, column=2, value=e['employee_name']).font = font_normal
        ws_ex2.cell(row=r_i, column=3, value=e['cargo']).font = font_normal
        ws_ex2.cell(row=r_i, column=4, value=e['departamento']).font = font_normal

        c_val = ws_ex2.cell(row=r_i, column=5, value=e['valor_auditado'])
        c_val.number_format = currency_format
        c_val.alignment = Alignment(horizontal="right")

        c_st = ws_ex2.cell(row=r_i, column=6, value=e['status'])
        c_st.alignment = Alignment(horizontal="center")
        if mode == 'beneficios':
            c_st.fill = fill_error
            c_st.font = font_error
        else:
            c_st.fill = fill_ok
            c_st.font = font_ok

        for col_b in range(1, 7):
            ws_ex2.cell(row=r_i, column=col_b).border = thin_border
            if r_i % 2 == 1 and mode != 'beneficios':
                ws_ex2.cell(row=r_i, column=col_b).fill = fill_zebra

    # Autoajuste de largura de colunas para todas as abas
    for ws in wb.worksheets:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = cell.value
                if val:
                    val_str = str(val)
                    if len(val_str) > max_len and len(val_str) < 60:
                        max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # Ajuste especial na aba de resumo
    ws_resumo.column_dimensions["B"].width = 38
    ws_resumo.column_dimensions["C"].width = 24
    ws_resumo.column_dimensions["D"].width = 50

    wb.save(output_path)
    return output_path


# ==============================================================================
# 4. GERADOR DE BASES DEMONSTRATIVAS DE EXEMPLO (FOLHA E BENEFÍCIOS)
# ==============================================================================

def create_demo_files(output_dir: str, mode: str = 'folha'):
    """
    Gera arquivos de teste com cenários realistas de RH:
    - Modo 'folha': Folha Mês Anterior vs Folha Mês Atual
    - Modo 'beneficios': Ativos RH vs Fatura Plano de Saúde
    """
    os.makedirs(output_dir, exist_ok=True)

    if mode == 'beneficios':
        # 1. Base Ativos no RH
        ativos_rh = [
            {"CPF": "048.123.456-78", "Matricula": "RH-101", "Nome_Colaborador": "Ana Beatriz Silva", "Departamento": "Tecnologia", "Status": "Ativo", "Plano": "Ouro"},
            {"CPF": "123.456.789-01", "Matricula": "RH-102", "Nome_Colaborador": "Carlos Eduardo Lima", "Departamento": "Financeiro", "Status": "Ativo", "Plano": "Prata"},
            {"CPF": "234.567.890-12", "Matricula": "RH-103", "Nome_Colaborador": "Juliana Santos Rocha", "Departamento": "Comercial", "Status": "Ativo", "Plano": "Ouro"},
            {"CPF": "345.678.901-23", "Matricula": "RH-104", "Nome_Colaborador": "Marcos Vinicius Alves", "Departamento": "Operações", "Status": "Ativo", "Plano": "Bronze"},
            {"CPF": "456.789.012-34", "Matricula": "RH-105", "Nome_Colaborador": "Fernanda Ribeiro Costa", "Departamento": "Marketing", "Status": "Ativo", "Plano": "Prata"},
            {"CPF": "567.890.123-45", "Matricula": "RH-106", "Nome_Colaborador": "Roberto Souza Mendes", "Departamento": "Jurídico", "Status": "Ativo", "Plano": "Ouro"},
            {"CPF": "678.901.234-56", "Matricula": "RH-107", "Nome_Colaborador": "Camila Oliveira Nunes", "Departamento": "RH", "Status": "Ativo", "Plano": "Prata"},
            {"CPF": "789.012.345-67", "Matricula": "RH-108", "Nome_Colaborador": "Lucas Pereira Braga", "Departamento": "Tecnologia", "Status": "Ativo", "Plano": "Bronze"}
        ]
        df_rh = pd.DataFrame(ativos_rh)

        # 2. Fatura da Operadora de Plano de Saúde (com "funcionário fantasma" e colaborador não faturado)
        fatura_saude = [
            # Ativos faturados corretamente
            {"CPF_Titular": "04812345678", "Nome_Segurado": "Ana Beatriz Silva", "Plano_Saude": "Ouro", "Mensalidade_R$": 850.00, "Coparticipacao_R$": 120.00, "Valor_Total": 970.00},
            {"CPF_Titular": "12345678901", "Nome_Segurado": "Carlos Eduardo Lima", "Plano_Saude": "Prata", "Mensalidade_R$": 620.00, "Coparticipacao_R$": 45.00, "Valor_Total": 665.00},
            {"CPF_Titular": "23456789012", "Nome_Segurado": "Juliana Santos Rocha", "Plano_Saude": "Ouro", "Mensalidade_R$": 850.00, "Coparticipacao_R$": 0.00, "Valor_Total": 850.00},
            {"CPF_Titular": "34567890123", "Nome_Segurado": "Marcos Vinicius Alves", "Plano_Saude": "Bronze", "Mensalidade_R$": 410.00, "Coparticipacao_R$": 60.00, "Valor_Total": 470.00},
            {"CPF_Titular": "45678901234", "Nome_Segurado": "Fernanda Ribeiro Costa", "Plano_Saude": "Prata", "Mensalidade_R$": 620.00, "Coparticipacao_R$": 80.00, "Valor_Total": 700.00},
            {"CPF_Titular": "56789012345", "Nome_Segurado": "Roberto Souza Mendes", "Plano_Saude": "Ouro", "Mensalidade_R$": 850.00, "Coparticipacao_R$": 150.00, "Valor_Total": 1000.00},
            # FUNCIONÁRIO FANTASMA 1: Já foi demitido, mas continua cobrando na fatura!
            {"CPF_Titular": "99988877766", "Nome_Segurado": "Joao Pedro Demitido", "Plano_Saude": "Ouro", "Mensalidade_R$": 850.00, "Coparticipacao_R$": 210.00, "Valor_Total": 1060.00},
            # FUNCIONÁRIO FANTASMA 2: Cobrança avulsa não identificada
            {"CPF_Titular": "88877766655", "Nome_Segurado": "Mariana Fantasma Ramos", "Plano_Saude": "Prata", "Mensalidade_R$": 620.00, "Coparticipacao_R$": 0.00, "Valor_Total": 620.00}
            # Camila Oliveira (67890123456) e Lucas Pereira (78901234567) estão no RH mas não constam na fatura!
        ]
        df_fatura = pd.DataFrame(fatura_saude)

        path_ref = os.path.join(output_dir, "Base_Ativos_RH_Exemplo.xlsx")
        path_audit = os.path.join(output_dir, "Fatura_Plano_Saude_Exemplo.xlsx")

        df_rh.to_excel(path_ref, index=False)
        df_fatura.to_excel(path_audit, index=False)

        return path_ref, "Base_Ativos_RH_Exemplo.xlsx", path_audit, "Fatura_Plano_Saude_Exemplo.xlsx"

    else:
        # Modo 'folha': Comparativo Mensal de Folha
        folha_anterior = [
            # CPF com zero à esquerda intencional sem pontuação para testar .zfill(11)
            {"CPF": "04812345678", "Matricula": "1001", "Nome_Colaborador": "Ana Beatriz Silva", "Cargo": "Analista Pleno", "Departamento": "Tecnologia", "Salario_Base": 5200.00, "Adicionais": 300.00, "Total_Proventos": 5500.00, "Descontos": 750.00, "Salario_Liquido": 4750.00},
            {"CPF": "123.456.789-01", "Matricula": "1002", "Nome_Colaborador": "Carlos Eduardo Lima", "Cargo": "Coordenador", "Departamento": "Financeiro", "Salario_Base": 7800.00, "Adicionais": 500.00, "Total_Proventos": 8300.00, "Descontos": 1250.00, "Salario_Liquido": 7050.00},
            {"CPF": "234.567.890-12", "Matricula": "1003", "Nome_Colaborador": "Juliana Santos Rocha", "Cargo": "Analista Jr", "Departamento": "Comercial", "Salario_Base": 3500.00, "Adicionais": 0.00, "Total_Proventos": 3500.00, "Descontos": 420.00, "Salario_Liquido": 3080.00},
            {"CPF": "345.678.901-23", "Matricula": "1004", "Nome_Colaborador": "Marcos Vinicius Alves", "Cargo": "Assistente", "Departamento": "Operações", "Salario_Base": 2800.00, "Adicionais": 250.00, "Total_Proventos": 3050.00, "Descontos": 310.00, "Salario_Liquido": 2740.00},
            {"CPF": "456.789.012-34", "Matricula": "1005", "Nome_Colaborador": "Fernanda Ribeiro Costa", "Cargo": "Especialista", "Departamento": "Marketing", "Salario_Base": 9200.00, "Adicionais": 0.00, "Total_Proventos": 9200.00, "Descontos": 1600.00, "Salario_Liquido": 7600.00},
            {"CPF": "567.890.123-45", "Matricula": "1006", "Nome_Colaborador": "Roberto Souza Mendes", "Cargo": "Advogado", "Departamento": "Jurídico", "Salario_Base": 6500.00, "Adicionais": 0.00, "Total_Proventos": 6500.00, "Descontos": 980.00, "Salario_Liquido": 5520.00},
            {"CPF": "678.901.234-56", "Matricula": "1007", "Nome_Colaborador": "Camila Oliveira Nunes", "Cargo": "Gerente RH", "Departamento": "RH", "Salario_Base": 11500.00, "Adicionais": 1000.00, "Total_Proventos": 12500.00, "Descontos": 2200.00, "Salario_Liquido": 10300.00},
            # Colaborador que será desligado na folha atual:
            {"CPF": "789.012.345-67", "Matricula": "1008", "Nome_Colaborador": "Lucas Pereira Braga", "Cargo": "Desenvolvedor", "Departamento": "Tecnologia", "Salario_Base": 6000.00, "Adicionais": 400.00, "Total_Proventos": 6400.00, "Descontos": 950.00, "Salario_Liquido": 5450.00}
        ]
        df_ant = pd.DataFrame(folha_anterior)

        folha_atual = [
            # 1. Ana Beatriz: 100% idêntica
            {"CPF": "048.123.456-78", "Matricula": "1001", "Nome_Colaborador": "Ana Beatriz Silva", "Cargo": "Analista Pleno", "Departamento": "Tecnologia", "Salario_Base": 5200.00, "Adicionais": 300.00, "Total_Proventos": 5500.00, "Descontos": 750.00, "Salario_Liquido": 4750.00},
            # 2. Carlos Eduardo: Aumento Salarial Atípico (+30%, Promoção para Gerente)
            {"CPF": "123.456.789-01", "Matricula": "1002", "Nome_Colaborador": "Carlos Eduardo Lima", "Cargo": "Gerente Financeiro", "Departamento": "Financeiro", "Salario_Base": 10140.00, "Adicionais": 500.00, "Total_Proventos": 10640.00, "Descontos": 1750.00, "Salario_Liquido": 8890.00},
            # 3. Juliana: Promoção salarial (+40% de 3500 para 4900) e mudança de cargo
            {"CPF": "234.567.890-12", "Matricula": "1003", "Nome_Colaborador": "Juliana Santos Rocha", "Cargo": "Analista Pleno", "Departamento": "Comercial", "Salario_Base": 4900.00, "Adicionais": 0.00, "Total_Proventos": 4900.00, "Descontos": 680.00, "Salario_Liquido": 4220.00},
            # 4. Marcos Vinicius: Aumento de Adicional de Horas Extras (oscilação moderada)
            {"CPF": "345.678.901-23", "Matricula": "1004", "Nome_Colaborador": "Marcos Vinicius Alves", "Cargo": "Assistente", "Departamento": "Operações", "Salario_Base": 2800.00, "Adicionais": 850.00, "Total_Proventos": 3650.00, "Descontos": 390.00, "Salario_Liquido": 3260.00},
            # 5. Fernanda: 100% idêntica
            {"CPF": "456.789.012-34", "Matricula": "1005", "Nome_Colaborador": "Fernanda Ribeiro Costa", "Cargo": "Especialista", "Departamento": "Marketing", "Salario_Base": 9200.00, "Adicionais": 0.00, "Total_Proventos": 9200.00, "Descontos": 1600.00, "Salario_Liquido": 7600.00},
            # 6. Roberto: Redução por falta ou licença (-15%)
            {"CPF": "567.890.123-45", "Matricula": "1006", "Nome_Colaborador": "Roberto Souza Mendes", "Cargo": "Advogado", "Departamento": "Jurídico", "Salario_Base": 5525.00, "Adicionais": 0.00, "Total_Proventos": 5525.00, "Descontos": 850.00, "Salario_Liquido": 4675.00},
            # 7. Camila: 100% idêntica
            {"CPF": "678.901.234-56", "Matricula": "1007", "Nome_Colaborador": "Camila Oliveira Nunes", "Cargo": "Gerente RH", "Departamento": "RH", "Salario_Base": 11500.00, "Adicionais": 1000.00, "Total_Proventos": 12500.00, "Descontos": 2200.00, "Salario_Liquido": 10300.00},
            # NOVO ADMITIDO: Não existia no mês anterior!
            {"CPF": "901.234.567-89", "Matricula": "1009", "Nome_Colaborador": "Gabriel Mendes Fontes", "Cargo": "Analista DevOps", "Departamento": "Tecnologia", "Salario_Base": 7000.00, "Adicionais": 0.00, "Total_Proventos": 7000.00, "Descontos": 1050.00, "Salario_Liquido": 5950.00}
        ]
        df_atual = pd.DataFrame(folha_atual)

        path_ref = os.path.join(output_dir, "Folha_Pagamento_Mes_Anterior.xlsx")
        path_audit = os.path.join(output_dir, "Folha_Pagamento_Mes_Atual.xlsx")

        df_ant.to_excel(path_ref, index=False)
        df_atual.to_excel(path_audit, index=False)

        return path_ref, "Folha_Pagamento_Mes_Anterior.xlsx", path_audit, "Folha_Pagamento_Mes_Atual.xlsx"
